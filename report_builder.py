from __future__ import annotations

from io import BytesIO

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _add_table(doc: Document, df: pd.DataFrame, columns: list[str], limit: int = 15) -> None:
    view = df.loc[:, [c for c in columns if c in df.columns]].head(limit).copy()
    table = doc.add_table(rows=1, cols=len(view.columns))
    table.style = "Table Grid"
    for idx, col in enumerate(view.columns):
        table.rows[0].cells[idx].text = str(col)
    for _, row in view.iterrows():
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = "—" if pd.isna(value) else (f"{value:.2f}" if isinstance(value, float) else str(value))


def build_report(fund: str, period: str, summary: dict, changes: pd.DataFrame,
                 themes: pd.DataFrame, peers: pd.DataFrame, notes: str, source: str,
                 industry_supplement: pd.DataFrame | None = None, holding_period: str = "未提供") -> bytes:
    doc = Document()
    title = doc.add_heading("基金分析報告", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"{fund}　{period}").bold = True

    doc.add_heading("投資結論", level=1)
    doc.add_paragraph(summary["conclusion"])
    doc.add_heading("基金規模與風險", level=1)
    metrics = pd.DataFrame({"項目": list(summary["metrics"].keys()), "數值": list(summary["metrics"].values())})
    _add_table(doc, metrics, ["項目", "數值"], 20)

    doc.add_heading("區間持股變化", level=1)
    _add_table(doc, changes, ["ticker", "name", "theme", "動作", "期初權重", "期末權重", "權重變化"], 20)
    doc.add_heading("投資題材", level=1)
    _add_table(doc, themes, ["theme", "期初權重", "期末權重", "權重變化"], 15)
    if industry_supplement is not None and not industry_supplement.empty:
        doc.add_heading("基金產業配置補充", level=1)
        _add_table(doc, industry_supplement, ["配置類型", "產業／題材", "權重", "資料日期", "個股明細"], 15)
        doc.add_paragraph("此為基金整體產業配置，不是單一個股；個股明細未揭露，因此不納入持股變化與獲利／虧損歸因。")

    doc.add_heading("持股損益歸因", level=1)
    doc.add_paragraph(f"淨值分析期間：{period}；持股市場報酬來源期間：{holding_period}。兩者可能不同，個股估計貢獻無法完整解釋基金損益。")
    valid = changes.dropna(subset=["估計貢獻"])
    winners = valid.loc[valid["估計貢獻"] > 0].sort_values("估計貢獻", ascending=False)
    losers = valid.loc[valid["估計貢獻"] < 0].sort_values("估計貢獻")
    doc.add_heading("獲利貢獻", level=2)
    if winners.empty:
        doc.add_paragraph("可計算的揭露持股中，沒有正估計貢獻。")
    else:
        _add_table(doc, winners, ["ticker", "name", "theme", "區間報酬", "估計貢獻"], 10)
    doc.add_heading("虧損拖累", level=2)
    if losers.empty:
        doc.add_paragraph("可計算的揭露持股中，沒有負估計貢獻；未揭露持股、現金、費用與交易影響仍可能造成基金回撤。")
    else:
        _add_table(doc, losers, ["ticker", "name", "theme", "區間報酬", "估計貢獻", "動作"], 10)
        for _, row in losers.head(10).iterrows():
            average_weight = (row["期初權重"] + row["期末權重"]) / 2
            doc.add_paragraph(
                f"{row['name']}：股價區間下跌 {abs(row['區間報酬']):.2f}%，平均權重 {average_weight:.2f}%，"
                f"估計拖累 {abs(row['估計貢獻']):.2f} 個百分點；持股{row['動作']}。"
            )
    doc.add_paragraph("虧損說明只根據個股價格和持股權重，不推測未經證實的事件原因。")

    doc.add_heading("同類基金比較", level=1)
    _add_table(doc, peers, ["基金", "區間報酬", "年化波動", "最大回撤", "Sharpe"], 15)
    doc.add_heading("補充說明", level=1)
    doc.add_paragraph(notes or "無")
    doc.add_paragraph(f"資料來源：{source or '使用者上傳資料'}")
    doc.add_paragraph("本報告為量化整理工具，不構成投資建議。持股貢獻採平均權重乘區間報酬估算，與基金正式歸因可能不同。")

    for section in doc.sections:
        section.header.paragraphs[0].text = "基金分析報告產生器"
    for style in doc.styles:
        if hasattr(style, "font"):
            style.font.name = "Arial"
            if style.name == "Normal":
                style.font.size = Pt(10.5)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()

